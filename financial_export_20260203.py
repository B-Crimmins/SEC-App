#!/usr/bin/env python3
"""
SEC Financial Data Export Tool - Improved Version (2026-02-03)

This script fetches financial data from the SEC for a given company ticker and year,
categorizes the data by financial statement type (Income Statement, Balance Sheet,
Cash Flow Statement), and exports the results to an Excel spreadsheet.

IMPROVEMENTS OVER ORIGINAL:
- Layer 1: Word-boundary-aware keyword matching (fixes "sales" in "AvailableForSale")
- Layer 2: Complete taxonomy map with calculation hierarchy for extension roll-ups
- Layer 3: Static mappings for deprecated tags still used by companies
- Layer 4: Balance type + periodType discriminator for accurate categorization

Usage:
    python financial_export_20260203.py

    Then enter the ticker symbol (e.g., MSFT) and year (e.g., 2024) when prompted.
"""

import sys
import re
import requests
import json
from collections import defaultdict
import zipfile
import io
import xml.etree.ElementTree as ET

# Excel export
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)

# SEC requires a User-Agent header
HEADERS = {
    'User-Agent': 'SEC-Financial-Export support@example.com',
}

INCLUDE_DEPRECATED = False

# =============================================================================
# LAYER 3: DEPRECATED TAG MAPPINGS
# These tags are no longer in current FASB taxonomy but companies still use them
# =============================================================================
DEPRECATED_TAG_MAPPINGS = {
    # Balance Sheet items incorrectly matched by keywords
    'AvailableForSaleSecurities': 'Balance Sheet',
    'AvailableForSaleSecuritiesCurrent': 'Balance Sheet',
    'AvailableForSaleSecuritiesNoncurrent': 'Balance Sheet',
    'AvailableForSaleSecuritiesAccumulatedGrossUnrealizedGainBeforeTax': 'Balance Sheet',
    'AvailableForSaleSecuritiesAccumulatedGrossUnrealizedLossBeforeTax': 'Balance Sheet',
    'AvailableForSaleSecuritiesAmortizedCost': 'Balance Sheet',
    'AvailableForSaleSecuritiesContinuousUnrealizedLossPosition12MonthsOrLongerAccumulatedLoss': 'Balance Sheet',
    'AvailableForSaleSecuritiesContinuousUnrealizedLossPositionAccumulatedLoss': 'Balance Sheet',
    'AvailableForSaleSecuritiesContinuousUnrealizedLossPositionFairValue': 'Balance Sheet',
    'AvailableForSaleSecuritiesContinuousUnrealizedLossPositionLessThan12MonthsAccumulatedLoss': 'Balance Sheet',
    'AvailableForSaleSecuritiesContinuousUnrealizedLossPositionLessThanTwelveMonthsFairValue': 'Balance Sheet',
    'AvailableForSaleSecuritiesContinuousUnrealizedLossPositionTwelveMonthsOrLongerFairValue': 'Balance Sheet',
    'DeferredTaxAssetsNetCurrent': 'Balance Sheet',
    'DeferredTaxAssetsNetNoncurrent': 'Balance Sheet',
    'DeferredTaxAssetsLiabilitiesNetCurrent': 'Balance Sheet',
    'DeferredTaxAssetsLiabilitiesNetNoncurrent': 'Balance Sheet',
    'DeferredTaxLiabilitiesNoncurrent': 'Balance Sheet',
    'CostMethodInvestments': 'Balance Sheet',
    'HeldToMaturitySecuritiesAmortizedCostBeforeOtherThanTemporaryImpairment': 'Balance Sheet',
    'HeldToMaturitySecuritiesContinuousUnrealizedLossPositionLessThanTwelveMonthsFairValue': 'Balance Sheet',
    'HeldToMaturitySecuritiesContinuousUnrealizedLossPositionTwelveMonthsOrLongerFairValue': 'Balance Sheet',

    # Cash Flow items incorrectly categorized
    'EffectOfExchangeRateOnCashAndCashEquivalents': 'Cash Flow Statement',
    'CashAndCashEquivalentsPeriodIncreaseDecrease': 'Cash Flow Statement',
    'ExcessTaxBenefitFromShareBasedCompensationFinancingActivities': 'Cash Flow Statement',
    'ExcessTaxBenefitFromShareBasedCompensationOperatingActivities': 'Cash Flow Statement',
    'ProceedsFromSaleOfAvailableForSaleSecurities': 'Cash Flow Statement',

    # Income Statement items
    'AvailableForSaleSecuritiesGrossRealizedGains': 'Income Statement',
    'AvailableForSaleSecuritiesGrossRealizedLosses': 'Income Statement',
    'ImpairmentOfInvestments': 'Income Statement',
    'LeaseAndRentalExpense': 'Income Statement',
    'RecognitionOfDeferredRevenue': 'Income Statement',
    'DeferredRevenueRevenueRecognized1': 'Income Statement',

    # Disclosure items that should be excluded
    'BusinessAcquisitionProFormaEarningsPerShareDiluted': 'Disclosure/Notes',
    'BusinessAcquisitionsProFormaNetIncomeLoss': 'Disclosure/Notes',
    'BusinessAcquisitionsProFormaRevenue': 'Disclosure/Notes',
    'MinorityInterestOwnershipPercentageByParent': 'Disclosure/Notes',
    'RestructuringAndRelatedCostExpectedNumberOfPositionsEliminated': 'Disclosure/Notes',
    'RestructuringAndRelatedCostNumberOfPositionsEliminatedInceptionToDate': 'Disclosure/Notes',
}


def get_cik_from_ticker(ticker: str) -> tuple:
    """
    Look up the CIK number from a ticker symbol using SEC's company_tickers.json

    Returns:
        tuple: (CIK padded to 10 digits, company name) or (None, None) if not found
    """
    print(f"Looking up CIK for ticker: {ticker.upper()}...")

    url = "https://www.sec.gov/files/company_tickers.json"

    try:
        response = requests.get(url, headers=HEADERS, timeout=30)
        response.raise_for_status()
        tickers_data = response.json()

        ticker_upper = ticker.upper()

        for entry in tickers_data.values():
            if entry.get('ticker', '').upper() == ticker_upper:
                cik = str(entry.get('cik_str', '')).zfill(10)
                company_name = entry.get('title', 'Unknown Company')
                print(f"Found: {company_name} (CIK: {cik})")
                return cik, company_name

        print(f"Ticker '{ticker}' not found in SEC database.")
        return None, None

    except Exception as e:
        print(f"Error looking up ticker: {e}")
        return None, None


def download_company_facts(cik: str) -> dict:
    """Download company facts from SEC EDGAR API"""
    print(f"Downloading company facts from SEC...")

    url = f"https://data.sec.gov/api/xbrl/companyfacts/CIK{cik}.json"

    try:
        response = requests.get(url, headers=HEADERS, timeout=60)
        response.raise_for_status()
        return response.json()
    except Exception as e:
        print(f"Error downloading company facts: {e}")
        return None


# =============================================================================
# LAYER 2: COMPLETE TAXONOMY MAP WITH CALCULATION HIERARCHY
# =============================================================================
def download_and_parse_taxonomy(year: int) -> dict:
    """
    Download FASB taxonomy and parse:
    - Presentation linkbases for statement categorization
    - Calculation linkbases for parent-child hierarchy
    - Schema files for balance type information (Layer 4)
    """
    print(f"\nDownloading FASB {year} taxonomy...")

    taxonomy_url = f"https://xbrl.fasb.org/us-gaap/{year}/us-gaap-{year}.zip"

    try:
        response = requests.get(taxonomy_url, timeout=120)
        response.raise_for_status()
        print("Taxonomy downloaded, extracting...")

        zip_file = zipfile.ZipFile(io.BytesIO(response.content))

        # Initialize result structure
        taxonomy_data = {
            'tag_to_category': {},      # Direct tag -> statement mapping
            'parent_child': {},          # Calculation hierarchy for roll-ups
            'balance_types': {},         # Balance type info for Layer 4
        }

        # Get all relevant files
        all_files = zip_file.namelist()
        linkbase_files = [f for f in all_files if (f.endswith(f'-pre-{year}.xml') or f.endswith(f'-cal-{year}.xml'))]
        schema_files = [f for f in all_files if f.endswith('.xsd')]

        print(f"Found {len(linkbase_files)} linkbase files, parsing...")

        # =================================================================
        # PART 1: Parse linkbases for tag categorization
        # =================================================================
        def get_category_from_filepath(filepath):
            lower_path = filepath.lower()

            # Statement files (primary financial statements)
            if '/stm/' in lower_path:
                if any(x in lower_path for x in ['soi', 'soc']):
                    return 'Income Statement'
                elif 'sfp' in lower_path:
                    return 'Balance Sheet'
                elif 'scf' in lower_path:
                    return 'Cash Flow Statement'
                elif any(x in lower_path for x in ['sheci', 'spc']):
                    return 'Statement of Equity'

            # Disclosure files - map to appropriate statement when possible
            if '/dis/' in lower_path:
                # Income Statement related disclosures
                if any(x in lower_path for x in [
                    'revenue', '-sr-', 'inctax', 'eps', 'disops', 'earnings',
                    'expense', 'ebit', 'compensation', 'crcrb', 'crcsbp',
                ]):
                    return 'Income Statement'

                # Balance Sheet related disclosures
                if any(x in lower_path for x in [
                    'debt', 'inv-', 'inventory', 'ppe', 'goodwill', 'intangible',
                    'leas', 'receivable', 'payable', 'securities', 'derivative',
                    'loan', 'deposit', 'property', '-re-', 'othliab',
                    'ides', 'bsoff', 'diha', 'fifvd', 'cc-', 'cecl',
                ]):
                    return 'Balance Sheet'

                if any(x in lower_path for x in ['scf', 'cashflow']):
                    return 'Cash Flow Statement'

                if any(x in lower_path for x in ['-se-', 'sharehold', 'stockholder']):
                    return 'Statement of Equity'

            return None

        # Parse linkbase files for tag -> category mapping
        for i, file_path in enumerate(linkbase_files):
            try:
                if (i + 1) % 50 == 0:
                    print(f"  Processed {i + 1}/{len(linkbase_files)} linkbase files...")

                category = get_category_from_filepath(file_path)
                if not category:
                    continue

                xml_content = zip_file.read(file_path)
                root = ET.fromstring(xml_content)

                ns = {
                    'link': 'http://www.xbrl.org/2003/linkbase',
                    'xlink': 'http://www.w3.org/1999/xlink'
                }

                # Get locators for tag names
                locators = root.findall('.//{http://www.xbrl.org/2003/linkbase}loc')

                for loc in locators:
                    href = loc.get('{http://www.w3.org/1999/xlink}href', '')

                    if '#us-gaap_' in href:
                        tag_name = href.split('#us-gaap_')[1]
                        full_tag = f'us-gaap:{tag_name}'

                        # Only set if not already mapped (first match wins for statements)
                        if full_tag not in taxonomy_data['tag_to_category']:
                            taxonomy_data['tag_to_category'][full_tag] = category

                # =============================================================
                # PART 2: Parse calculation arcs for parent-child relationships
                # =============================================================
                if '-cal-' in file_path:
                    calc_links = root.findall('.//link:calculationLink', ns)

                    for calc_link in calc_links:
                        # Build locator label -> concept mapping
                        concept_labels = {}
                        locs = calc_link.findall('link:loc', ns)
                        for loc in locs:
                            label = loc.get('{http://www.w3.org/1999/xlink}label')
                            href = loc.get('{http://www.w3.org/1999/xlink}href', '')
                            if '#us-gaap_' in href:
                                concept = href.split('#us-gaap_')[1]
                                concept_labels[label] = concept

                        # Parse calculation arcs
                        arcs = calc_link.findall('link:calculationArc', ns)
                        for arc in arcs:
                            from_label = arc.get('{http://www.w3.org/1999/xlink}from')
                            to_label = arc.get('{http://www.w3.org/1999/xlink}to')

                            parent = concept_labels.get(from_label)
                            child = concept_labels.get(to_label)

                            if parent and child:
                                if child not in taxonomy_data['parent_child']:
                                    taxonomy_data['parent_child'][child] = parent

            except Exception:
                pass

        # =================================================================
        # PART 3: Parse schema for balance types (Layer 4)
        # =================================================================
        print("Parsing schema for balance types...")
        main_schema = [f for f in schema_files if f.endswith(f'us-gaap-{year}.xsd')]

        if main_schema:
            try:
                xml_content = zip_file.read(main_schema[0])
                root = ET.fromstring(xml_content)

                ns = {'xs': 'http://www.w3.org/2001/XMLSchema'}

                elements = root.findall('.//xs:element', ns)
                for elem in elements:
                    name = elem.get('name', '')
                    balance = elem.get('{http://www.xbrl.org/2003/instance}balance')
                    period_type = elem.get('{http://www.xbrl.org/2003/instance}periodType')

                    if name:
                        taxonomy_data['balance_types'][name] = {
                            'balance': balance,
                            'periodType': period_type
                        }

                print(f"  Extracted balance types for {len(taxonomy_data['balance_types'])} concepts")
            except Exception as e:
                print(f"  Warning: Could not parse schema for balance types: {e}")

        print(f"Taxonomy parsing complete:")
        print(f"  - {len(taxonomy_data['tag_to_category'])} tags categorized")
        print(f"  - {len(taxonomy_data['parent_child'])} parent-child relationships")
        print(f"  - {len(taxonomy_data['balance_types'])} balance types")

        return taxonomy_data

    except Exception as e:
        print(f"Error downloading taxonomy: {e}")
        print("Falling back to keyword-based categorization only.")
        return {'tag_to_category': {}, 'parent_child': {}, 'balance_types': {}}


# =============================================================================
# LAYER 1: WORD-BOUNDARY-AWARE KEYWORD MATCHING
# =============================================================================
def split_camel_case(tag_name: str) -> set:
    """Split camelCase tag name into individual words"""
    # Split on uppercase letters, keeping the letter with the following word
    words = re.findall(r'[A-Z][a-z]*|[a-z]+', tag_name)
    return {w.lower() for w in words}


def categorize_by_keywords(tag_name: str) -> str:
    """
    Improved keyword categorization using word boundary awareness.

    Key improvements:
    - Splits camelCase into words to avoid substring matches
    - Checks Balance Sheet indicators BEFORE Income Statement
    - More precise keyword sets
    """
    # Split camelCase into individual words
    words = split_camel_case(tag_name)
    tag_lower = tag_name.lower()

    # Disclosure/Note items - skip these first
    disclosure_keywords = {'disclosure', 'textblock', 'policy', 'table', 'abstract',
                           'lineitem', 'domain', 'member', 'axis'}
    if words & disclosure_keywords:
        return 'Disclosure/Notes'

    # =================================================================
    # IMPORTANT: Check Balance Sheet FIRST for items with position words
    # This prevents "DeferredTaxAssetsNetCurrent" from matching "tax" first
    # =================================================================
    balance_position_words = {'current', 'noncurrent', 'net', 'gross'}
    balance_type_words = {'assets', 'asset', 'liabilities', 'liability',
                          'receivable', 'receivables', 'payable', 'payables',
                          'equity', 'deficit', 'stock', 'shares'}

    # If has position word AND balance sheet type word -> Balance Sheet
    if (words & balance_position_words) and (words & balance_type_words):
        return 'Balance Sheet'

    # Cash Flow Statement - check specific patterns
    # Use full-word matching for most, but allow compound words
    cashflow_exact_words = {'cashflow', 'cashflows'}
    cashflow_phrase_patterns = [
        'operatingactivities', 'investingactivities', 'financingactivities',
        'periodincreasedecrease', 'effectofexchangerate'
    ]
    cashflow_words = {'proceeds', 'repayment', 'repayments', 'issuance',
                      'depreciation', 'amortization'}

    if words & cashflow_exact_words:
        return 'Cash Flow Statement'
    if any(pattern in tag_lower for pattern in cashflow_phrase_patterns):
        return 'Cash Flow Statement'
    if words & cashflow_words:
        # Additional check: these could also be disclosure items
        if 'payment' in words and not ('cash' in words or 'debt' in words):
            pass  # Don't categorize yet
        else:
            return 'Cash Flow Statement'

    # Income Statement - use word boundaries
    income_words = {'revenue', 'revenues', 'income', 'expense', 'expenses',
                    'earnings', 'profit', 'loss', 'losses',
                    'margin', 'ebit', 'ebitda'}

    # Specific patterns that indicate Income Statement
    income_patterns = ['costof', 'grossprofit', 'operatingincome', 'netincome',
                       'comprehensiveincome', 'earningspershare']

    if words & income_words:
        # But NOT if it's clearly a balance sheet item
        if not (words & {'deferred', 'accumulated', 'retained'}):
            return 'Income Statement'

    if any(pattern in tag_lower for pattern in income_patterns):
        return 'Income Statement'

    # Balance Sheet - remaining items
    balance_words = {'cash', 'inventory', 'inventories', 'property', 'debt',
                     'goodwill', 'intangible', 'intangibles', 'investment',
                     'investments', 'securities', 'derivative', 'derivatives',
                     'lease', 'leases', 'deposit', 'deposits', 'loan', 'loans',
                     'land', 'building', 'buildings', 'equipment',
                     'accumulated', 'allowance', 'deferred', 'accrued',
                     'retained', 'treasury', 'capital', 'surplus'}

    if words & balance_words:
        return 'Balance Sheet'

    return 'Other'


# =============================================================================
# LAYER 4: BALANCE TYPE DISCRIMINATOR
# =============================================================================
def categorize_with_balance_type(tag_name: str, keyword_result: str, balance_types: dict) -> tuple:
    """
    Use balance type and periodType to validate/override keyword categorization.

    Returns:
        tuple: (final_category, confidence_level, override_reason)

    Decision matrix:
        debit + instant = Balance Sheet (Asset)
        credit + instant = Balance Sheet (Liability/Equity)
        debit + duration = Income Statement (Expense/Loss)
        credit + duration = Income Statement (Revenue/Gain)
        no balance + duration = Often Cash Flow
    """
    # Strip us-gaap: prefix if present
    clean_tag = tag_name.replace('us-gaap:', '')

    info = balance_types.get(clean_tag, {})
    balance = info.get('balance')
    period_type = info.get('periodType')

    # If no balance type info, return keyword result with medium confidence
    if not period_type:
        return (keyword_result, 'medium', None)

    # Apply decision matrix
    if period_type == 'instant':
        # Instant = point-in-time = Balance Sheet item
        if keyword_result != 'Balance Sheet' and keyword_result not in ['Disclosure/Notes', 'Other']:
            return ('Balance Sheet', 'high', f'periodType=instant overrides {keyword_result}')
        return (keyword_result, 'high', None)

    elif period_type == 'duration':
        # Duration = over-time = Income Statement or Cash Flow
        if keyword_result == 'Balance Sheet':
            # This is likely wrong - duration items are flows, not stocks
            if balance in ['debit', 'credit']:
                return ('Income Statement', 'high', f'periodType=duration overrides Balance Sheet')
            else:
                return ('Cash Flow Statement', 'medium', f'periodType=duration + no balance suggests Cash Flow')
        return (keyword_result, 'high', None)

    return (keyword_result, 'medium', None)


# =============================================================================
# LAYER 2 CONTINUED: EXTENSION ROLL-UP
# =============================================================================
def find_parent_category(tag_name: str, taxonomy_data: dict, max_depth: int = 5) -> str:
    """
    Walk up the calculation hierarchy to find a categorized parent.
    This handles extension tags that roll up to standard GAAP parents.
    """
    clean_tag = tag_name.replace('us-gaap:', '')
    parent_child = taxonomy_data.get('parent_child', {})
    tag_to_category = taxonomy_data.get('tag_to_category', {})

    current = clean_tag
    visited = set()

    for _ in range(max_depth):
        if current in visited:
            break
        visited.add(current)

        parent = parent_child.get(current)
        if not parent:
            break

        # Check if parent has a category
        parent_full = f'us-gaap:{parent}'
        if parent_full in tag_to_category:
            return tag_to_category[parent_full]

        current = parent

    return None


# =============================================================================
# MAIN CATEGORIZATION FUNCTION
# =============================================================================
def categorize_tag(tag_name: str, taxonomy_data: dict) -> tuple:
    """
    Multi-layer categorization combining all strategies.

    Returns:
        tuple: (category, confidence, method)
    """
    full_tag = tag_name if tag_name.startswith('us-gaap:') else f'us-gaap:{tag_name}'
    clean_tag = tag_name.replace('us-gaap:', '')

    # Layer 3: Check deprecated mappings first
    if clean_tag in DEPRECATED_TAG_MAPPINGS:
        return (DEPRECATED_TAG_MAPPINGS[clean_tag], 'high', 'deprecated_mapping')

    # Layer 2a: Check taxonomy direct mapping
    if full_tag in taxonomy_data.get('tag_to_category', {}):
        category = taxonomy_data['tag_to_category'][full_tag]
        return (category, 'high', 'taxonomy_direct')

    # Layer 2b: Try parent roll-up
    parent_category = find_parent_category(full_tag, taxonomy_data)
    if parent_category:
        return (parent_category, 'medium-high', 'taxonomy_rollup')

    # Layer 1: Keyword-based categorization
    keyword_result = categorize_by_keywords(clean_tag)

    # Layer 4: Validate/override with balance type
    balance_types = taxonomy_data.get('balance_types', {})
    final_category, confidence, override_reason = categorize_with_balance_type(
        clean_tag, keyword_result, balance_types
    )

    method = 'keyword'
    if override_reason:
        method = f'keyword+balance_override'

    return (final_category, confidence, method)


def process_company_data(company_data: dict, taxonomy_data: dict, target_year: int, form_type: str = '10-K') -> tuple:
    """
    Process company facts and categorize by statement type.

    Returns:
        tuple: (categorized_data, provenance_data)
            - categorized_data: dict of items by statement category
            - provenance_data: dict with filing source statistics and restatements
    """
    print(f"\nProcessing financial data for year {target_year} ({form_type} filings)...")

    categorized = defaultdict(list)

    # Provenance tracking
    provenance = {
        'target_year': target_year,
        'form_type': form_type,
        'original_filing_count': 0,
        'comparative_filing_count': 0,
        'restatements': [],  # Items where value changed between filings
        'filing_sources': defaultdict(int),  # Count by filing date
        'accession_numbers': set(),  # All accession numbers used
    }

    stats = {
        'total_tags': 0,
        'taxonomy_direct': 0,
        'taxonomy_rollup': 0,
        'deprecated_mapping': 0,
        'keyword': 0,
        'keyword_balance_override': 0,
        'skipped_deprecated': 0,
        'skipped_no_year_data': 0,
        'by_confidence': defaultdict(int)
    }

    if 'facts' not in company_data or 'us-gaap' not in company_data['facts']:
        print("No US GAAP facts found in company data.")
        return categorized, provenance

    us_gaap_facts = company_data['facts']['us-gaap']
    stats['total_tags'] = len(us_gaap_facts)
    print(f"Found {stats['total_tags']} total US GAAP tags")

    for tag_name, tag_data in us_gaap_facts.items():
        full_tag = f"us-gaap:{tag_name}"
        label = tag_data.get('label', '') or tag_name

        # Skip deprecated tags if configured
        if not INCLUDE_DEPRECATED and 'Deprecated' in label:
            stats['skipped_deprecated'] += 1
            continue

        # Multi-layer categorization
        category, confidence, method = categorize_tag(tag_name, taxonomy_data)

        # Track statistics
        if 'balance_override' in method:
            stats['keyword_balance_override'] += 1
        elif method == 'taxonomy_direct':
            stats['taxonomy_direct'] += 1
        elif method == 'taxonomy_rollup':
            stats['taxonomy_rollup'] += 1
        elif method == 'deprecated_mapping':
            stats['deprecated_mapping'] += 1
        else:
            stats['keyword'] += 1

        stats['by_confidence'][confidence] += 1

        # Skip non-statement categories
        if category in ['Disclosure/Notes', 'Other', 'Unknown']:
            continue

        # Filter by target year and get ALL matching values for provenance tracking
        if 'units' not in tag_data:
            stats['skipped_no_year_data'] += 1
            continue

        for unit, values in tag_data['units'].items():
            # Get all values matching target year and form type
            # Filter to entries where 'end' date is in target year (actual FY data)
            year_values = [v for v in values
                          if v.get('end', '').startswith(str(target_year)) and
                             v.get('form') == form_type]

            if not year_values:
                continue

            # Group by end date to properly detect restatements
            # A restatement is when the SAME period (same end date) has different values
            by_end_date = defaultdict(list)
            for v in year_values:
                end_date = v.get('end', '')
                by_end_date[end_date].append(v)

            # Use the latest end date (typically the fiscal year end)
            latest_end_date = max(by_end_date.keys())
            entries_for_period = by_end_date[latest_end_date]

            # Sort by filing date (oldest first for comparison)
            sorted_values = sorted(entries_for_period, key=lambda x: x.get('filed', ''))

            # Get original (first) and latest (last) filings for this period
            original_entry = sorted_values[0]
            latest_entry = sorted_values[-1]

            original_value = original_entry.get('val')
            latest_value = latest_entry.get('val')
            original_filed = original_entry.get('filed', 'Unknown')
            latest_filed = latest_entry.get('filed', 'Unknown')
            original_accn = original_entry.get('accn', '')
            latest_accn = latest_entry.get('accn', '')
            original_fy = original_entry.get('fy', target_year)
            latest_fy = latest_entry.get('fy', target_year)

            # Determine source type based on which filing the data came from
            # If the latest filing's FY is greater than target year, it's comparative data
            if latest_fy > target_year:
                is_comparative = True
                source = f"Comparative (FY{latest_fy} {form_type})"
            elif len(sorted_values) > 1 and latest_filed != original_filed:
                # Multiple filings for same period, using the latest one
                is_comparative = True
                source = f"Revised (filed {latest_filed})"
            else:
                is_comparative = False
                source = f"Original (FY{target_year} {form_type})"

            # Check for restatement (value changed between filings FOR THE SAME PERIOD)
            is_restated = False
            restatement_delta = None
            restatement_pct = None

            # Only flag as restatement if same end date has different values
            if len(sorted_values) > 1 and original_value != latest_value:
                is_restated = True
                if original_value and original_value != 0:
                    restatement_delta = latest_value - original_value
                    restatement_pct = (restatement_delta / original_value) * 100

                    # Track restatement for provenance report
                    provenance['restatements'].append({
                        'tag': full_tag,
                        'label': label,
                        'category': category,
                        'original_value': original_value,
                        'original_filed': original_filed,
                        'original_accn': original_accn,
                        'latest_value': latest_value,
                        'latest_filed': latest_filed,
                        'latest_accn': latest_accn,
                        'delta': restatement_delta,
                        'pct_change': restatement_pct,
                        'unit': unit,
                        'period_end': latest_end_date
                    })

            # Track provenance statistics
            if is_comparative:
                provenance['comparative_filing_count'] += 1
            else:
                provenance['original_filing_count'] += 1

            provenance['filing_sources'][latest_filed] += 1
            provenance['accession_numbers'].add(latest_accn)
            if original_accn != latest_accn:
                provenance['accession_numbers'].add(original_accn)

            # Add to categorized data with full provenance info
            categorized[category].append({
                'tag': full_tag,
                'label': label,
                'value': latest_value,
                'unit': unit,
                'filed': latest_filed,
                'confidence': confidence,
                'method': method,
                # New provenance fields
                'source': source,
                'is_restated': is_restated,
                'original_value': original_value if is_restated else None,
                'original_filed': original_filed if is_restated else None,
                'restatement_delta': restatement_delta,
                'restatement_pct': restatement_pct,
                'accession': latest_accn,
            })
            break  # Only process first unit type

        else:
            stats['skipped_no_year_data'] += 1

    # Print statistics
    print(f"\nCategorization Statistics:")
    print(f"  Taxonomy direct matches: {stats['taxonomy_direct']}")
    print(f"  Taxonomy roll-ups: {stats['taxonomy_rollup']}")
    print(f"  Deprecated mappings: {stats['deprecated_mapping']}")
    print(f"  Keyword matches: {stats['keyword']}")
    print(f"  Balance type overrides: {stats['keyword_balance_override']}")
    print(f"  Skipped deprecated: {stats['skipped_deprecated']}")
    print(f"  Skipped (no {target_year} data): {stats['skipped_no_year_data']}")
    print(f"\nConfidence breakdown:")
    for conf, count in sorted(stats['by_confidence'].items()):
        print(f"  {conf}: {count}")

    # Print provenance summary
    print(f"\nData Provenance:")
    print(f"  Items from original FY{target_year} filing: {provenance['original_filing_count']}")
    print(f"  Items from comparative filings: {provenance['comparative_filing_count']}")
    print(f"  Restatements detected: {len(provenance['restatements'])}")

    return categorized, provenance


def export_to_excel(categorized_data: dict, provenance_data: dict, ticker: str, company_name: str, year: int, form_type: str, output_file: str):
    """Export categorized data to Excel spreadsheet with provenance tracking"""
    print(f"\nExporting to Excel: {output_file}")

    wb = Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    category_fills = {
        'Income Statement': PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        'Balance Sheet': PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
        'Cash Flow Statement': PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        'Statement of Equity': PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    }

    # Confidence-based fills
    confidence_fills = {
        'high': PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        'medium-high': PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        'medium': PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    }

    # Restatement highlight fill
    restated_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    restated_yes_fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Desired order for sheets
    sheet_order = ['Income Statement', 'Balance Sheet', 'Cash Flow Statement', 'Statement of Equity']

    first_sheet = True
    for category in sheet_order:
        if category not in categorized_data or not categorized_data[category]:
            continue

        items = categorized_data[category]

        if first_sheet:
            ws = wb.active
            ws.title = category[:31]  # Excel sheet name limit
            first_sheet = False
        else:
            ws = wb.create_sheet(title=category[:31])

        # Header info
        ws['A1'] = f"{company_name} ({ticker.upper()})"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"Fiscal Year: {year}"
        ws['A2'].font = Font(bold=True, size=12)
        ws['A3'] = f"Statement Type: {category}"
        ws['A3'].font = Font(bold=True, size=12)
        ws.merge_cells('A1:K1')
        ws.merge_cells('A2:K2')
        ws.merge_cells('A3:K3')

        # Column headers (added Source, Restated, Original Value columns)
        headers = ['XBRL Tag', 'Label', 'Value', 'Unit', 'Filing Date', 'Source', 'Restated', 'Original Value', 'Delta', 'Confidence', 'Method']
        header_row = 5

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # Data rows
        category_fill = category_fills.get(category, PatternFill())

        for row_idx, item in enumerate(items, header_row + 1):
            # Basic columns
            ws.cell(row=row_idx, column=1, value=item['tag']).border = thin_border
            ws.cell(row=row_idx, column=2, value=item['label']).border = thin_border

            value_cell = ws.cell(row=row_idx, column=3, value=item['value'])
            value_cell.border = thin_border
            if isinstance(item['value'], (int, float)):
                value_cell.number_format = '#,##0'

            ws.cell(row=row_idx, column=4, value=item['unit']).border = thin_border
            ws.cell(row=row_idx, column=5, value=item['filed']).border = thin_border

            # Source column (new)
            source_cell = ws.cell(row=row_idx, column=6, value=item.get('source', 'Unknown'))
            source_cell.border = thin_border

            # Restated column (new)
            is_restated = item.get('is_restated', False)
            restated_cell = ws.cell(row=row_idx, column=7, value='Yes' if is_restated else 'No')
            restated_cell.border = thin_border
            if is_restated:
                restated_cell.fill = restated_yes_fill
                restated_cell.font = Font(bold=True, color="FFFFFF")

            # Original Value column (new) - only show if restated
            orig_val_cell = ws.cell(row=row_idx, column=8)
            orig_val_cell.border = thin_border
            if is_restated and item.get('original_value') is not None:
                orig_val_cell.value = item['original_value']
                if isinstance(item['original_value'], (int, float)):
                    orig_val_cell.number_format = '#,##0'
                orig_val_cell.fill = restated_fill
            else:
                orig_val_cell.value = '-'

            # Delta column (new) - show change with percentage
            delta_cell = ws.cell(row=row_idx, column=9)
            delta_cell.border = thin_border
            if is_restated and item.get('restatement_delta') is not None:
                delta = item['restatement_delta']
                pct = item.get('restatement_pct', 0)
                delta_cell.value = f"{delta:+,.0f} ({pct:+.2f}%)"
                delta_cell.fill = restated_fill
            else:
                delta_cell.value = '-'

            # Confidence column
            conf_cell = ws.cell(row=row_idx, column=10, value=item.get('confidence', 'unknown'))
            conf_cell.border = thin_border
            conf_fill = confidence_fills.get(item.get('confidence'), PatternFill())
            conf_cell.fill = conf_fill

            # Method column
            ws.cell(row=row_idx, column=11, value=item.get('method', 'unknown')).border = thin_border

            # Apply category fill to main data columns (not the new provenance columns)
            for col in range(1, 6):
                ws.cell(row=row_idx, column=col).fill = category_fill

        # Auto-adjust column widths
        column_widths = [45, 55, 18, 12, 12, 28, 10, 18, 22, 12, 22]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

    # ==========================================================================
    # Summary sheet
    # ==========================================================================
    ws_summary = wb.create_sheet(title="Summary", index=0)
    ws_summary['A1'] = f"SEC Financial Data Export (Improved)"
    ws_summary['A1'].font = Font(bold=True, size=16)
    ws_summary['A3'] = f"Company: {company_name}"
    ws_summary['A4'] = f"Ticker: {ticker.upper()}"
    ws_summary['A5'] = f"Fiscal Year: {year}"
    ws_summary['A6'] = f"Form Type: {form_type}"
    ws_summary['A8'] = "Data by Statement Type:"
    ws_summary['A8'].font = Font(bold=True)

    row = 9
    total_items = 0
    for category in sheet_order:
        if category in categorized_data:
            count = len(categorized_data[category])
            total_items += count
            ws_summary[f'A{row}'] = f"  {category}:"
            ws_summary[f'B{row}'] = f"{count} items"
            row += 1

    ws_summary[f'A{row + 1}'] = "Total:"
    ws_summary[f'A{row + 1}'].font = Font(bold=True)
    ws_summary[f'B{row + 1}'] = f"{total_items} items"
    ws_summary[f'B{row + 1}'].font = Font(bold=True)

    # Add provenance summary to Summary sheet
    row += 4
    ws_summary[f'A{row}'] = "Data Provenance Summary:"
    ws_summary[f'A{row}'].font = Font(bold=True)
    row += 1
    ws_summary[f'A{row}'] = f"  Items from original FY{year} filing:"
    ws_summary[f'B{row}'] = provenance_data.get('original_filing_count', 0)
    row += 1
    ws_summary[f'A{row}'] = f"  Items from comparative filings:"
    ws_summary[f'B{row}'] = provenance_data.get('comparative_filing_count', 0)
    row += 1
    ws_summary[f'A{row}'] = f"  Restatements detected:"
    ws_summary[f'B{row}'] = len(provenance_data.get('restatements', []))

    # Add confidence legend
    row += 3
    ws_summary[f'A{row}'] = "Column Legend:"
    ws_summary[f'A{row}'].font = Font(bold=True)
    row += 1
    ws_summary[f'A{row}'] = "  Source: Indicates which filing the data came from"
    row += 1
    ws_summary[f'A{row}'] = "    - 'Original': Data from the FY's own 10-K filing"
    row += 1
    ws_summary[f'A{row}'] = "    - 'Comparative': Data from a subsequent year's 10-K (as prior-year comparison)"
    row += 1
    ws_summary[f'A{row}'] = "  Restated: 'Yes' if the value changed between original and subsequent filings"
    row += 1
    ws_summary[f'A{row}'] = "  Original Value: The value from the original filing (only shown if restated)"
    row += 1
    ws_summary[f'A{row}'] = "  Delta: The change amount and percentage (only shown if restated)"

    row += 2
    ws_summary[f'A{row}'] = "Confidence Legend:"
    ws_summary[f'A{row}'].font = Font(bold=True)
    row += 1
    ws_summary[f'A{row}'] = "  high: Direct taxonomy match or balance type confirmed"
    row += 1
    ws_summary[f'A{row}'] = "  medium-high: Parent roll-up from calculation hierarchy"
    row += 1
    ws_summary[f'A{row}'] = "  medium: Keyword match or balance type inference"

    ws_summary.column_dimensions['A'].width = 55
    ws_summary.column_dimensions['B'].width = 15

    # ==========================================================================
    # Data Provenance sheet (new)
    # ==========================================================================
    ws_prov = wb.create_sheet(title="Data Provenance")

    ws_prov['A1'] = "Data Provenance Report"
    ws_prov['A1'].font = Font(bold=True, size=16)

    ws_prov['A3'] = f"Target Fiscal Year: {year}"
    ws_prov['A4'] = f"Form Type: {form_type}"

    # Filing sources breakdown
    ws_prov['A6'] = "Filing Sources:"
    ws_prov['A6'].font = Font(bold=True)

    row = 7
    for filing_date, count in sorted(provenance_data.get('filing_sources', {}).items()):
        ws_prov[f'A{row}'] = f"  {filing_date}:"
        ws_prov[f'B{row}'] = f"{count} items"
        row += 1

    # Accession numbers for audit trail
    row += 1
    ws_prov[f'A{row}'] = "SEC Accession Numbers (for audit trail):"
    ws_prov[f'A{row}'].font = Font(bold=True)
    row += 1
    for accn in sorted(provenance_data.get('accession_numbers', set())):
        if accn:
            ws_prov[f'A{row}'] = f"  {accn}"
            row += 1

    # Restatements section
    row += 2
    ws_prov[f'A{row}'] = "Restatements Detected:"
    ws_prov[f'A{row}'].font = Font(bold=True)
    row += 1

    restatements = provenance_data.get('restatements', [])
    if restatements:
        # Headers for restatement table
        restatement_headers = ['XBRL Tag', 'Label', 'Category', 'Original Value', 'Original Filed',
                               'Latest Value', 'Latest Filed', 'Delta', '% Change']
        for col, header in enumerate(restatement_headers, 1):
            cell = ws_prov.cell(row=row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
        row += 1

        # Restatement data
        for restatement in restatements:
            ws_prov.cell(row=row, column=1, value=restatement['tag']).border = thin_border
            ws_prov.cell(row=row, column=2, value=restatement['label']).border = thin_border
            ws_prov.cell(row=row, column=3, value=restatement['category']).border = thin_border

            orig_cell = ws_prov.cell(row=row, column=4, value=restatement['original_value'])
            orig_cell.border = thin_border
            if isinstance(restatement['original_value'], (int, float)):
                orig_cell.number_format = '#,##0'

            ws_prov.cell(row=row, column=5, value=restatement['original_filed']).border = thin_border

            latest_cell = ws_prov.cell(row=row, column=6, value=restatement['latest_value'])
            latest_cell.border = thin_border
            if isinstance(restatement['latest_value'], (int, float)):
                latest_cell.number_format = '#,##0'

            ws_prov.cell(row=row, column=7, value=restatement['latest_filed']).border = thin_border

            delta_cell = ws_prov.cell(row=row, column=8, value=restatement.get('delta', 0))
            delta_cell.border = thin_border
            if isinstance(restatement.get('delta'), (int, float)):
                delta_cell.number_format = '+#,##0;-#,##0'

            pct_cell = ws_prov.cell(row=row, column=9, value=restatement.get('pct_change', 0))
            pct_cell.border = thin_border
            pct_cell.number_format = '+0.00%;-0.00%'

            # Highlight the row
            for col in range(1, 10):
                ws_prov.cell(row=row, column=col).fill = restated_fill

            row += 1
    else:
        ws_prov[f'A{row}'] = "  No restatements detected - all values match between original and subsequent filings."
        row += 1

    # Explanation
    row += 2
    ws_prov[f'A{row}'] = "Notes:"
    ws_prov[f'A{row}'].font = Font(bold=True)
    row += 1
    ws_prov[f'A{row}'] = "- 'Comparative' data comes from subsequent years' 10-K filings where FY{} appears as prior-year comparison.".format(year)
    row += 1
    ws_prov[f'A{row}'] = "- Using the latest filing ensures you have the most recent audited values (including any corrections)."
    row += 1
    ws_prov[f'A{row}'] = "- Restatements indicate the company revised previously reported values in a subsequent filing."

    # Adjust column widths
    prov_widths = [45, 50, 18, 18, 12, 18, 12, 18, 12]
    for col, width in enumerate(prov_widths, 1):
        ws_prov.column_dimensions[get_column_letter(col)].width = width

    # Save workbook
    wb.save(output_file)
    print(f"Excel file saved successfully!")
    print(f"Total items exported: {total_items}")


def main():
    print("=" * 70)
    print("SEC Financial Data Export Tool - IMPROVED VERSION (2026-02-03)")
    print("=" * 70)
    print()
    print("Improvements in this version:")
    print("  - Layer 1: Word-boundary-aware keyword matching")
    print("  - Layer 2: Taxonomy hierarchy for extension roll-ups")
    print("  - Layer 3: Deprecated tag mappings")
    print("  - Layer 4: Balance type + periodType validation")
    print()

    # Get user input
    ticker = input("Enter ticker symbol (e.g., MSFT, AAPL): ").strip()
    if not ticker:
        print("Error: Ticker symbol is required.")
        sys.exit(1)

    year_input = input("Enter fiscal year (e.g., 2024): ").strip()
    try:
        year = int(year_input)
        if year < 2000 or year > 2030:
            raise ValueError("Year out of reasonable range")
    except ValueError:
        print("Error: Please enter a valid year (e.g., 2024).")
        sys.exit(1)

    print("\nForm types:")
    print("  1. 10-K (Annual Report)")
    print("  2. 10-Q (Quarterly Report)")
    form_choice = input("Select form type [1]: ").strip()

    if form_choice == '2':
        form_type = '10-Q'
    else:
        form_type = '10-K'

    print(f"\nSelected: {form_type}")
    print()

    # Step 1: Look up CIK from ticker
    cik, company_name = get_cik_from_ticker(ticker)
    if not cik:
        print(f"Could not find CIK for ticker '{ticker}'. Please verify the ticker symbol.")
        sys.exit(1)

    # Step 2: Download company facts
    company_data = download_company_facts(cik)
    if not company_data:
        print("Failed to download company data.")
        sys.exit(1)

    # Step 3: Download and parse taxonomy (enhanced)
    taxonomy_data = download_and_parse_taxonomy(year)

    # Step 4: Process and categorize data (multi-layer) with provenance tracking
    categorized_data, provenance_data = process_company_data(company_data, taxonomy_data, year, form_type)

    if not any(categorized_data.values()):
        print(f"\nNo financial data found for {ticker.upper()} in year {year} ({form_type}).")
        sys.exit(1)

    # Step 5: Export to Excel with provenance information
    output_file = f"{ticker.upper()}_{year}_{form_type}_financial_data_improved.xlsx"
    export_to_excel(categorized_data, provenance_data, ticker, company_name, year, form_type, output_file)

    print()
    print("=" * 70)
    print(f"Export complete: {output_file}")
    print("=" * 70)


if __name__ == "__main__":
    main()
